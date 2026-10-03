# -*- coding: utf-8 -*-
"""核验重生成的 docx 是否已含 A1/A2 新内容、且不含陈旧「未执行」表述。
用法：python results/g_build/check_a1_in_docx.py   （须系统 Python 3.11 + python-docx）
"""
import os
import docx

DOCS = {
    r"E:\learnflow\docs\研究报告_中文版.docx": ["1,274", "0.963", "0.892"],
    r"E:\learnflow\docs\研究报告_韩文版.docx": ["1,274", "0.963", "0.892"],
    r"E:\learnflow\docs\研究计划_中文版.docx": ["1,274", "0.963"],
    r"E:\learnflow\docs\研究计划_韩文版.docx": ["1,274", "0.963"],
}
STALE = [
    "还须再补能力校正判据",          # 研究报告 中
    "두 가지 추가(**미수행**)",       # 研究计划 韩
    "(**미수행**)",                   # 研究计划 韩（宽松）
    "（**未执行**）",                 # 研究计划 中
    "⚠️ 未执行，已在 M1 声明",        # 研究报告 中
    "⚠️ 미실행, M1에서",              # 研究报告 韩
]


def full_text(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def main():
    all_ok = True
    for path, needles in DOCS.items():
        if not os.path.exists(path):
            print("MISSING:", path)
            all_ok = False
            continue
        txt = full_text(path)
        missing = [n for n in needles if n not in txt]
        stale = [s for s in STALE if s in txt]
        ok = (not missing) and (not stale)
        all_ok = all_ok and ok
        print("%-28s %s | 缺失新内容=%s | 检出旧表述=%s"
              % (os.path.basename(path), "OK" if ok else "FAIL", missing, stale))
    print("\n总体:", "PASS 全部通过" if all_ok else "FAIL 需修正")
    return 0 if all_ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
