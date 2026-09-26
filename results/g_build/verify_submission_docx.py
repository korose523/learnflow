# -*- coding: utf-8 -*-
"""G 任务：验证 10-20 提交包 docx 的时效性与关键内容。

校验：
  1. 核心属性 author / last_modified_by 非空（审阅 §2.5 署名质疑）；
  2. 产物 mtime >= 源 md mtime（未过期）；
  3. 关键更正内容确实写入（§4.5 被试计划 / §7.1 期刊口径核验 / 因果侧时点声明）。
"""
import os
from docx import Document

CHECKS = [
    (r"E:\learnflow\docs\研究计划_韩文版.docx",
     r"E:\learnflow\docs\研究计划与报告\研究计划_韩文版.md",
     ["학과에서 정하는 소정의 연구 실적", "영산대학교 IRB"]),
    (r"E:\learnflow\docs\研究计划_中文版.docx",
     r"E:\learnflow\docs\研究计划与报告\研究计划_中文版.md",
     ["该学科自行规定", "Youngsan 大学 IRB"]),
    (r"E:\learnflow\docs\研究报告_韩文版.docx",
     r"E:\learnflow\docs\研究计划与报告\研究报告_韩文版.md",
     ["1,326"]),
    (r"E:\learnflow\docs\研究报告_中文版.docx",
     r"E:\learnflow\docs\研究计划与报告\研究报告_中文版.md",
     ["1,326"]),
    (r"E:\learnflow\docs\因果侧识别策略.docx",
     r"E:\learnflow\docs\因果侧识别策略.md",
     ["未接触"]),
]


def full_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def main():
    all_ok = True
    for docx, md, needles in CHECKS:
        if not os.path.exists(docx):
            print("✗ 缺失产物:", os.path.basename(docx))
            all_ok = False
            continue
        d = Document(docx)
        cp = d.core_properties
        txt = full_text(d)
        fresh = os.path.getmtime(docx) >= os.path.getmtime(md)
        author_ok = bool((cp.author or "").strip()) and bool((cp.last_modified_by or "").strip())
        print("── %s" % os.path.basename(docx))
        print("   author=%r  last_modified_by=%r  title=%r" % (cp.author, cp.last_modified_by, cp.title))
        print("   时效性: 产物>=源 %s" % ("✅" if fresh else "❌ 过期"))
        print("   作者属性非空: %s" % ("✅" if author_ok else "❌ 空"))
        for n in needles:
            hit = n in txt
            print("   含 %r: %s" % (n, "✅" if hit else "❌ 未找到"))
            if not hit:
                all_ok = False
        if not (fresh and author_ok):
            all_ok = False
        print("")
    print("总体: %s" % ("✅ 全部通过" if all_ok else "❌ 存在未通过项"))
    return 0 if all_ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
