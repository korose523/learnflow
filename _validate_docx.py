# -*- coding: utf-8 -*-
"""校验生成的 docx：字体、标题、目录、关键数字。"""
import io, os, re
from docx import Document
from docx.oxml.ns import qn

BASE = r"E:\learnflow\docs"
FILES = [
    ("研究计划_中文版.docx", "zh"),
    ("研究计划_韩文版.docx", "kr"),
    ("研究报告_中文版.docx", "zh"),
    ("研究报告_韩文版.docx", "kr"),
]

def doc_eastasia(doc):
    ea = set()
    for p in doc.paragraphs:
        for r in p.runs:
            rpr = r._element.find(qn("w:rPr"))
            if rpr is None:
                continue
            rf = rpr.find(qn("w:rFonts"))
            if rf is not None:
                v = rf.get(qn("w:eastAsia"))
                if v:
                    ea.add(v)
    return ea

def count_headings(doc):
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    hs = 0
    toc = False
    title = ""
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") or p.style.name == "Title":
            hs += 1
        if p.text.strip() in ("目录", "목차"):
            toc = True
        if p.style.name == "Title":
            title = p.text.strip()
    return hs, toc, title

def has_text(doc, needle):
    for p in doc.paragraphs:
        if needle in p.text:
            return True
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                if needle in c.text:
                    return True
    return False

lines = []
for fn, lang in FILES:
    dp = os.path.join(BASE, fn)
    doc = Document(dp)
    ea = doc_eastasia(doc)
    hs, toc, title = count_headings(doc)
    ntables = len(doc.tables)
    # 关键数字抽查
    checks = {
        "ee1a49be5732": has_text(doc, "ee1a49be5732"),
        "0.5357": has_text(doc, "0.5357"),
        "−0.1222(U+2212)": has_text(doc, "−0.1222"),
        "λ_cf(k)": has_text(doc, "λ_cf(k)") if "报告" in fn else True,
        "16,217,311": has_text(doc, "16,217,311") if "报告" in fn else True,
    }
    lines.append("=== %s (lang=%s) ===" % (fn, lang))
    lines.append("  段落=%d 表格=%d 标题段落(H/Title)=%d 目录插入=%s" % (len(doc.paragraphs), ntables, hs, toc))
    lines.append("  标题: %s" % (title[:40] if title else "(无)"))
    lines.append("  东亚字体集合: %s" % ", ".join(sorted(ea)))
    lines.append("  关键数字: " + ", ".join("%s=%s" % (k, v) for k, v in checks.items()))
    lines.append("")

io.open(r"E:\learnflow\_validate_docx.txt", "w", encoding="utf-8", newline="").write("\n".join(lines))
print("\n".join(lines))
